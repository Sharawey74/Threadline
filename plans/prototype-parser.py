#!/usr/bin/env python3
"""
Threadline — derived-model parser prototype.

Proves the "your files are the data" schema against the real Career folder.
READ-ONLY. This script never writes to any file under the career root.

Its job is not to be clever. Its job is to be *inspectable*: every extraction
records how confident it is and why, so ambiguity in the source files surfaces
as data instead of hiding as a silent wrong guess.
"""

from __future__ import annotations

import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

CAREER_ROOT = Path("/sessions/upbeat-wizardly-faraday/mnt/Career")
TASKS_FILE = "TASKS.md"
STUDY_DIR = "Study guided & notes"
MATRIX = "Roadmap tracking documents/roadmap_competency_matrix_v3.xlsx"

# Normalise the dashes and arrows that appear in the real file.
DASHES = {"—": "-", "–": "-", "→": "->"}


def norm(s: str) -> str:
    for a, b in DASHES.items():
        s = s.replace(a, b)
    return s


# ─────────────────────────────────────────────────────────────────────────────
# Model
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class HourMatch:
    value: float
    raw: str
    emphasis: str          # "italic" | "bold" | "tilde" | "plain"
    pos: int


@dataclass
class Task:
    line_no: int                    # 1-based, for safe write-back
    raw: str
    checked: bool
    section: str
    text: str                       # display text, markup stripped
    order: int | None = None        # leading "**1." if present
    hours: float | None = None
    hours_conf: str = "none"        # "high" | "low" | "none"
    pages: int | None = None
    pages_conf: str = "none"
    backticks: list[str] = field(default_factory=list)
    link_target: str | None = None  # resolved path, relative to career root
    link_kind: str = "none"         # "folder" | "file" | "ambiguous" | "unresolved"
    notes: list[str] = field(default_factory=list)


@dataclass
class Section:
    number: int | None
    title: str
    budget_hours: float | None
    line_no: int


@dataclass
class Topic:
    slug: str                       # "02 - Databases & Storage"
    order: int
    path: Path
    material: list[Path]

    @property
    def page_estimate(self) -> None:
        return None


# ─────────────────────────────────────────────────────────────────────────────
# Extractors
# ─────────────────────────────────────────────────────────────────────────────

HOUR_RE = re.compile(r"(?P<pre>\*{0,2}~?\(?)(?P<num>\d+(?:\.\d+)?)\s*h\b(?P<post>\)?\*{0,2})")
PAGE_RE = re.compile(r"(?P<pre>\(?)(?P<num>\d+)\s*(?:pp|pages)\b")
BACKTICK_RE = re.compile(r"`([^`]+)`")
CHECKBOX_RE = re.compile(r"^\s*-\s*\[(?P<mark>[ xX])\]\s*(?P<body>.*)$")
HEADING_RE = re.compile(r"^(?P<hashes>#{1,6})\s+(?P<body>.*)$")
SECTION_NUM_RE = re.compile(r"^(?P<num>\d+)\.\s+(?P<title>.*)$")
ORDER_RE = re.compile(r"^\*{0,2}(?P<num>\d+)\.\s")


def find_hours(text: str) -> list[HourMatch]:
    out = []
    for m in HOUR_RE.finditer(text):
        pre, post = m.group("pre"), m.group("post")
        if "~" in pre:
            emph = "tilde"
        elif pre.count("*") >= 2 or post.count("*") >= 2:
            emph = "bold"
        elif "*" in pre or "*" in post:
            emph = "italic"
        else:
            emph = "plain"
        out.append(HourMatch(float(m.group("num")), m.group(0), emph, m.start()))
    return out


def pick_hours(matches: list[HourMatch]) -> tuple[float | None, str, list[str]]:
    """Choose the hour figure that represents THIS task's allocation."""
    notes: list[str] = []
    if not matches:
        return None, "none", notes
    if len(matches) > 1:
        notes.append(
            "multiple hour figures: " + ", ".join(f"{m.raw.strip()}({m.emphasis})" for m in matches)
        )
    # Emphasised figures are the author's deliberate allocation markers.
    emphasised = [m for m in matches if m.emphasis in ("italic", "bold", "tilde")]
    if emphasised:
        chosen = emphasised[-1]
        return chosen.value, "high", notes
    return matches[-1].value, "low", notes


def find_pages(text: str) -> tuple[int | None, str, list[str]]:
    notes: list[str] = []
    ms = list(PAGE_RE.finditer(text))
    if not ms:
        return None, "none", notes
    if len(ms) > 1:
        notes.append("multiple page counts: " + ", ".join(m.group(0) for m in ms))
    # A page count inside parentheses usually describes a referenced file,
    # not the size of the task itself.
    free = [m for m in ms if not m.group("pre")]
    if free:
        return int(free[0].group("num")), "high", notes
    notes.append("page count only appears parenthesised - likely describes another file")
    return int(ms[0].group("num")), "low", notes


def strip_markup(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return re.sub(r"\s+", " ", s).strip()


# ─────────────────────────────────────────────────────────────────────────────
# Link resolution
# ─────────────────────────────────────────────────────────────────────────────

def resolve_link(root: Path, candidates: list[str]) -> tuple[str | None, str, list[str]]:
    """Resolve backtick spans against the real filesystem."""
    notes: list[str] = []
    hits: list[tuple[str, str]] = []
    for c in candidates:
        c_clean = c.strip()
        direct = root / c_clean
        if direct.is_dir():
            hits.append((c_clean, "folder"))
            continue
        if direct.is_file():
            hits.append((c_clean, "file"))
            continue
        # search by basename anywhere under the root
        found = [p for p in root.rglob("*") if p.name == c_clean]
        if len(found) == 1:
            rel = str(found[0].relative_to(root))
            hits.append((rel, "folder" if found[0].is_dir() else "file"))
        elif len(found) > 1:
            notes.append(f"`{c_clean}` matches {len(found)} paths")
        else:
            # try as a path fragment
            frag = [p for p in root.rglob("*") if str(p.relative_to(root)).endswith(c_clean)]
            if len(frag) == 1:
                hits.append((str(frag[0].relative_to(root)), "file" if frag[0].is_file() else "folder"))
            elif not frag:
                notes.append(f"`{c_clean}` does not exist on disk")

    if not hits:
        return None, "unresolved", notes
    folders = [h for h in hits if h[1] == "folder"]
    if len(hits) > 1:
        notes.append(f"{len(hits)} backtick links on one line")
    if len(folders) == 1:
        return folders[0][0], "folder", notes
    if len(hits) == 1:
        return hits[0][0], hits[0][1], notes
    return hits[0][0], "ambiguous", notes


# ─────────────────────────────────────────────────────────────────────────────
# TASKS.md
# ─────────────────────────────────────────────────────────────────────────────

def parse_tasks(root: Path) -> tuple[list[Section], list[Task], dict]:
    path = root / TASKS_FILE
    raw_lines = path.read_text(encoding="utf-8").splitlines()

    sections: list[Section] = []
    tasks: list[Task] = []
    meta: dict = {}
    current = "(preamble)"

    title_line = norm(raw_lines[0]) if raw_lines else ""
    if m := re.search(r"(\d{1,2} \w+ \d{4})\s*->\s*(\d{1,2} \w+ \d{4})", title_line):
        meta["window_start"], meta["window_end"] = m.group(1), m.group(2)

    for i, raw in enumerate(raw_lines, start=1):
        line = norm(raw)

        if hm := HEADING_RE.match(line):
            if len(hm.group("hashes")) == 2:
                body = hm.group("body")
                budget = None
                # "*(19h)*" and "*(Oct-Nov, 38h)*" both carry a budget.
                if b := re.search(r"\*\([^)]*?(\d+(?:\.\d+)?)\s*h\)\*", body):
                    budget = float(b.group(1))
                    body = body[: b.start()].strip()
                num = None
                if sm := SECTION_NUM_RE.match(body):
                    num, body = int(sm.group("num")), sm.group("title")
                sections.append(Section(num, strip_markup(body), budget, i))
                current = strip_markup(body)
            continue

        if cm := CHECKBOX_RE.match(line):
            body = cm.group("body")
            t = Task(
                line_no=i,
                raw=raw,
                checked=cm.group("mark").lower() == "x",
                section=current,
                text=strip_markup(body),
            )
            if om := ORDER_RE.match(body):
                t.order = int(om.group("num"))
            t.hours, t.hours_conf, hn = pick_hours(find_hours(body))
            t.notes += hn
            t.pages, t.pages_conf, pn = find_pages(body)
            t.notes += pn
            t.backticks = BACKTICK_RE.findall(body)
            if t.backticks:
                t.link_target, t.link_kind, ln = resolve_link(root, t.backticks)
                t.notes += ln
            tasks.append(t)

    # budget table in section 3
    budget: dict[str, float] = {}
    for raw in raw_lines:
        line = norm(raw).strip()
        if line.startswith("|") and line.count("|") >= 3:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if len(cells) == 2 and re.fullmatch(r"\*{0,2}\d+\*{0,2}", cells[1]):
                budget[strip_markup(cells[0])] = float(strip_markup(cells[1]))
    meta["budget"] = budget
    return sections, tasks, meta


# ─────────────────────────────────────────────────────────────────────────────
# Study topics
# ─────────────────────────────────────────────────────────────────────────────

TOPIC_RE = re.compile(r"^(?P<num>\d{2})\s*-\s*(?P<name>.+)$")
MATERIAL_EXT = {".pdf", ".md", ".docx", ".xlsx", ".pptx", ".html", ".txt"}


def scan_topics(root: Path) -> list[Topic]:
    base = root / STUDY_DIR
    topics = []
    if not base.is_dir():
        return topics
    for d in sorted(base.iterdir()):
        if not d.is_dir():
            continue
        m = TOPIC_RE.match(d.name)
        if not m:
            continue
        material = sorted(
            p for p in d.rglob("*")
            if p.is_file() and p.suffix.lower() in MATERIAL_EXT
        )
        topics.append(Topic(d.name, int(m.group("num")), d, material))
    return topics


# ─────────────────────────────────────────────────────────────────────────────
# Competency matrix
# ─────────────────────────────────────────────────────────────────────────────

def parse_matrix(root: Path) -> tuple[list[dict], list[str]]:
    problems: list[str] = []
    try:
        import openpyxl
    except ImportError:
        return [], ["openpyxl not installed"]
    p = root / MATRIX
    if not p.exists():
        return [], [f"{MATRIX} not found"]
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [str(c or "").strip() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        rows.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
    for name in wb.sheetnames[1:]:
        s = wb[name]
        filled = sum(1 for row in s.iter_rows(min_row=2, values_only=True)
                     for c in row[1:] if c not in (None, ""))
        if filled == 0:
            problems.append(f"sheet '{name}' has labels but no computed values")
    return rows, problems


# ─────────────────────────────────────────────────────────────────────────────
# Report
# ─────────────────────────────────────────────────────────────────────────────

def rule(title: str) -> None:
    print(f"\n{'=' * 78}\n{title}\n{'=' * 78}")


def main() -> int:
    root = CAREER_ROOT
    if not root.is_dir():
        print(f"career root not found: {root}", file=sys.stderr)
        return 1

    sections, tasks, meta = parse_tasks(root)
    topics = scan_topics(root)
    matrix, matrix_problems = parse_matrix(root)

    rule("PLAN")
    print(f"window          {meta.get('window_start')} -> {meta.get('window_end')}")
    print(f"sections        {len(sections)}")
    print(f"tasks           {len(tasks)}  ({sum(t.checked for t in tasks)} done)")
    if meta.get("budget"):
        total = meta["budget"].get("Total")
        print(f"budget table    {len(meta['budget'])} rows, total {total}h")

    rule("SECTIONS WITH AN HOUR BUDGET")
    for s in sections:
        if s.budget_hours:
            n = sum(1 for t in tasks if t.section == s.title)
            alloc = sum(t.hours or 0 for t in tasks if t.section == s.title)
            flag = "" if abs(alloc - s.budget_hours) < 0.01 else f"   <-- tasks sum to {alloc}h"
            print(f"  §{s.number:<3} {s.title:<28} {s.budget_hours:>6}h  {n:>2} tasks{flag}")

    rule("TASKS")
    for t in tasks:
        box = "x" if t.checked else " "
        h = f"{t.hours}h" if t.hours is not None else "--"
        conf = {"high": " ", "low": "?", "none": " "}[t.hours_conf]
        link = f"  -> {t.link_kind}:{t.link_target}" if t.link_target else ""
        print(f"  L{t.line_no:<4} [{box}] {h:>6}{conf} {t.text[:58]:<58}{link}")
        for n in t.notes:
            print(f"        ! {n}")

    rule("STUDY TOPICS ON DISK")
    total_material = 0
    for tp in topics:
        total_material += len(tp.material)
        linked = [t for t in tasks if t.link_target == f"{STUDY_DIR}/{tp.slug}"
                  or (t.link_target or "").endswith(tp.slug)]
        mark = f"linked to L{linked[0].line_no}" if linked else "NO TASK LINKS TO THIS"
        print(f"  {tp.slug:<42} {len(tp.material):>2} files   {mark}")
    print(f"\n  {len(topics)} topics, {total_material} material files")

    rule("COMPETENCY MATRIX")
    if matrix:
        import collections
        status = collections.Counter(str(r.get("Status")) for r in matrix)
        pillar = collections.Counter(str(r.get("Pillar")) for r in matrix)
        n = len(matrix)
        print(f"  {n} rows")
        for k, v in status.most_common():
            print(f"    {k:<14} {v:>4}   {v / n:>5.1%}")
        print()
        for k, v in pillar.most_common():
            done = sum(1 for r in matrix if r.get("Pillar") == k and r.get("Status") == "Done")
            print(f"    {k:<38} {done:>3}/{v:<3} done")
    for p in matrix_problems:
        print(f"  ! {p}")

    rule("AMBIGUITIES THE REAL FILES CONTAIN")
    flagged = [t for t in tasks if t.notes]
    if not flagged:
        print("  none")
    for t in flagged:
        print(f"  L{t.line_no}: {t.text[:64]}")
        for n in t.notes:
            print(f"      - {n}")
    print(f"\n  {len(flagged)} of {len(tasks)} task lines need a disambiguation rule")

    rule("RECONCILIATION  (does the plan add up?)")
    ok = True

    def check(label: str, got: float, want: float, unit: str = "h") -> None:
        nonlocal ok
        good = abs(got - want) < 0.01
        ok = ok and good
        print(f"  {'OK ' if good else 'X  '} {label:<44} {got:>7g}{unit} vs {want:g}{unit}")

    month_sections = [s for s in sections if s.budget_hours and s.number and s.number >= 8]
    month_total = 0.0
    for s in month_sections:
        alloc = sum(t.hours or 0 for t in tasks if t.section == s.title)
        month_total += s.budget_hours
        check(f"§{s.number} {s.title} tasks vs heading", alloc, s.budget_hours)
    check("month headings vs budget table total", month_total, meta["budget"].get("Total", 0))

    study_tasks = [t for t in tasks if t.link_kind == "folder" and t.order]
    check("study topic hours vs budget table row",
          sum(t.hours or 0 for t in study_tasks),
          meta["budget"].get("Study guide & notes - 9 topics, 396 pages", 0))
    print(f"  {'OK ' if len(study_tasks) == len(topics) else 'X  '} "
          f"{'study tasks vs topic folders on disk':<44} {len(study_tasks):>7} vs {len(topics)}")
    total_pp = sum(t.pages or 0 for t in study_tasks)
    check("study pages declared vs sum of topics", total_pp, 396, "pp")

    print(f"\n  {'the plan is internally consistent' if ok else 'DRIFT DETECTED'}")

    rule("WRITE-BACK TARGETS")
    print("  Ticking a box edits exactly one line. Byte offsets that would change:")
    for t in tasks[:4]:
        idx = t.raw.index("[")
        print(f"    L{t.line_no}  col {idx + 1}  '[{' ' if not t.checked else 'x'}]' -> '[x]'")
    print("  ...everything else in the file is untouched.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
